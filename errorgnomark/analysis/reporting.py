# File Path: errorgnomark/analysis/reporting.py

# ==============================================================
# [FINAL v16] Enhanced Reporting Framework
# --------------------------------------------------------------
# • Compatible with QST / MRB / QPT analysis
# • Automatic results_demo/<demo_name>/ directory
# • MRB Decay Plot + MRB Heatmap integrated
# ==============================================================

import os, io, inspect, itertools
import numpy as np
from typing import List, Dict, Any
import matplotlib.pyplot as plt
from mpl_toolkits.mplot3d import Axes3D  # noqa
from mpl_toolkits.axes_grid1 import make_axes_locatable

try:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.drawing.image import Image as OpenpyxlImage
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
except ImportError:
    print("Warning: 'openpyxl' not found, Excel features disabled.")
    Workbook = None
    ExcelReport = None

# --- Core ExperimentResult ---
from .result import ExperimentResult

# --- MRB Decay Function (fallback) ---
try:
    from .mrb import _mrb_decay_func
except (ImportError, ModuleNotFoundError):
    def _mrb_decay_func(m, A, p, B): return A * (p ** m) + B


# ==============================================================
# --- Plotting utilities for QPT/QST ---
# ==============================================================
def plot_density_matrix(rho: np.ndarray, title="Density Matrix"):
    if rho.ndim != 2 or rho.shape[0] != rho.shape[1]:
        raise ValueError("rho must be square")
    dim = rho.shape[0]
    nq = int(np.log2(dim))
    ticks = [format(i, f"0{nq}b") for i in range(dim)]
    fig = plt.figure(figsize=(14, 6))
    fig.suptitle(title, fontsize=16)
    x, y = np.meshgrid(np.arange(dim), np.arange(dim))
    maxz = max(np.max(np.abs(rho)), 1e-12)
    axr = fig.add_subplot(1, 2, 1, projection="3d")
    axr.bar3d(x.ravel(), y.ravel(), 0, 0.8, 0.8, np.real(rho).ravel())
    axr.set_title("Real Part"); axr.set_zlim(-maxz, maxz)
    axi = fig.add_subplot(1, 2, 2, projection="3d")
    axi.bar3d(x.ravel(), y.ravel(), 0, 0.8, 0.8, np.imag(rho).ravel())
    axi.set_title("Imag Part"); axi.set_zlim(-maxz, maxz)
    plt.tight_layout(rect=[0, 0.05, 1, 0.95])
    return fig


def plot_chi_matrix(chi: np.ndarray, title="Process (Chi) Matrix"):
    dim = chi.shape[0]
    n = int(round(np.log(dim) / np.log(4)))
    labels = [''.join(p) for p in itertools.product('IXYZ', repeat=n)]
    re, im = np.real(chi), np.imag(chi)
    vmax = np.max(np.abs(np.concatenate((re, im))))
    fig, axs = plt.subplots(1, 2, figsize=(12, 5))
    fig.suptitle(title, fontsize=16)
    for a, data, part in zip(axs, [re, im], ["Real", "Imaginary"]):
        im_ = a.imshow(data, cmap="RdBu_r", vmin=-vmax, vmax=vmax)
        a.set_xticks(range(dim)); a.set_yticks(range(dim))
        a.set_xticklabels(labels, rotation=90); a.set_yticklabels(labels)
        a.set_title(f"{part} Part")
        fig.colorbar(im_, ax=a, fraction=0.046, pad=0.04)
    plt.tight_layout(rect=[0, 0.05, 1, 0.95])
    return fig


# ==============================================================
def generate_report(results: List[Any]):
    print("\n" + "=" * 50)
    print(" " * 15 + "ANALYSIS REPORT")
    print("=" * 50)
    for i, r in enumerate(results):
        print(f"\n--- Result {i+1}: {getattr(r, 'name', f'Result_{i+1}')} ---")
        d = getattr(r, "data", {})
        if "process_fidelity" in d:
            print(f"  Process Fidelity: {d['process_fidelity']:.4f}")
        if "fidelity_with_ideal" in d:
            print(f"  State Fidelity: {d['fidelity_with_ideal']:.4f}")
        if "epc" in d:
            print(f"  EPC: {d['epc']:.3e} ± {d.get('epc_err',0):.2e}")
    print("\n" + "=" * 50)


# ==============================================================
# --- Excel Report Class (auto results_demo + MRB plots) ---
# ==============================================================
if Workbook:
    class ExcelReport:
        def __init__(self, output_dir: str = None):
            # auto-detect demo file
            demo = "unspecified_demo"
            for f in inspect.stack():
                if "demo_" in f.filename:
                    demo = os.path.splitext(os.path.basename(f.filename))[0]
                    break
            base = os.path.abspath(os.path.join(os.path.dirname(__file__), "../.."))
            default_dir = os.path.join(base, "results_demo", demo)
            self.output_dir = os.path.abspath(output_dir) if output_dir else default_dir
            os.makedirs(self.output_dir, exist_ok=True)

            self.workbook = Workbook()
            self.workbook.remove(self.workbook.active)
            self._header_font = Font(bold=True, color="FFFFFF")
            self._header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
            self._center_align = Alignment(horizontal="center", vertical="center")
            self._thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                                       top=Side(style='thin'), bottom=Side(style='thin'))

        # ---------- basics ----------
        def _apply_header_style(self, c):
            c.font, c.fill, c.alignment, c.border = (
                self._header_font, self._header_fill, self._center_align, self._thin_border
            )

        def _save_fig_to_sheet(self, fig: plt.Figure, sheet_name: str):
            ws = self.workbook.create_sheet(title=sheet_name)
            buf = io.BytesIO()
            fig.savefig(buf, format="png", dpi=200, bbox_inches="tight")
            plt.close(fig)
            buf.seek(0)
            ws.add_image(OpenpyxlImage(buf), "A1")

        # ---------- summary ----------
        def create_summary_sheet(self, analysis_results: List[ExperimentResult], experiment_params: Dict[str, Any]):
            """Creates parameter summary and key metrics table."""
            ws = self.workbook.create_sheet("Summary")
            ws.cell(1, 1, "Experiment Parameters").font = Font(bold=True, size=14)
            r = 2
            for k, v in experiment_params.items():
                ws.cell(r, 1, k).font = Font(bold=True)
                ws.cell(r, 2, str(v))
                r += 1
            r += 1
            ws.cell(r, 1, "Key Results").font = Font(bold=True, size=14)
            r += 1
            headers = ["Name", "Process Fidelity", "State Fidelity", "EPC"]
            for ci, h in enumerate(headers, 1):
                self._apply_header_style(ws.cell(r, ci, value=h))
            for res in analysis_results:
                r += 1
                ws.cell(r, 1, res.name).border = self._thin_border
                ws.cell(r, 2, res.data.get("process_fidelity")).border = self._thin_border
                ws.cell(r, 3, res.data.get("fidelity_with_ideal")).border = self._thin_border
                ws.cell(r, 4, res.data.get("epc")).border = self._thin_border
            self._auto_fit_columns(ws)
            return self

        # ---------- QPT / QST plots ----------
        def add_chi_matrix_plot(self, chi, sheet_name="Process Chi Matrix"):
            self._save_fig_to_sheet(plot_chi_matrix(chi, sheet_name), sheet_name)
            return self

        def add_density_matrix_plot(self, rho, sheet_name="Density Matrix"):
            self._save_fig_to_sheet(plot_density_matrix(rho, sheet_name), sheet_name)
            return self

        # ---------- MRB specific plots ----------
        def add_mrb_decay_plot(self, result: ExperimentResult):
            """Adds MRB decay curve plot."""
            fit = result.data.get("fit_params", {})
            raw = result.data.get("raw_data", {})
            group = result.data.get("group", "Unknown")
            if not raw:
                return self
            depths = np.array(sorted(raw.keys()))
            means = np.array([np.mean(raw[d]) for d in depths])
            errs = np.array([np.std(raw[d]) / max(len(raw[d]), 1)**0.5 for d in depths])
            fig, ax = plt.subplots(figsize=(8, 5))
            ax.errorbar(depths, means, yerr=errs, fmt="o", capsize=4, label="Mean Survival")
            if fit:
                A, p, B = fit.get("params", (1, 1, 0))
                xs = np.linspace(0, max(depths) * 1.05, 200)
                ax.plot(xs, _mrb_decay_func(xs, A, p, B), "r-", label=f"Fit p={p:.4f}")
            ax.set_xlabel("Depth (m)"); ax.set_ylabel("Survival Probability")
            ax.set_title(f"MRB Decay {group}")
            ax.legend(); ax.grid(True, alpha=0.6, ls="--")
            self._save_fig_to_sheet(fig, f"MRB Decay {group}")
            return self

        def add_mrb_heatmap(self, result: ExperimentResult):
            """Adds MRB polarization heatmap."""
            groups = result.data.get("qubit_groups", [])
            depths = result.data.get("depths", [])
            pol = np.array(result.data.get("avg_polarizations", []))
            if pol.size == 0:
                return self
            fig, ax = plt.subplots(figsize=(max(8, len(groups) * 1.5), max(5, len(depths) * 0.6)))
            im = ax.imshow(pol.T, cmap="Blues", vmin=0, vmax=1, aspect="auto")
            ax.set_xticks(range(len(groups)))
            ax.set_xticklabels([str(g) for g in groups], rotation=45, ha="right")
            ax.set_yticks(range(len(depths)))
            ax.set_yticklabels(depths)
            ax.set_xlabel("Qubit Groups"); ax.set_ylabel("Depths")
            ax.set_title("MRB Direct Polarization", pad=15)
            for i in range(pol.shape[0]):
                for j in range(pol.shape[1]):
                    v = pol[i, j]
                    ax.text(j, i, f"{v:.2f}", ha="center", va="center",
                            color="white" if v > 0.5 else "black", fontsize=8)
            divider = make_axes_locatable(ax)
            cax = divider.append_axes("right", size="5%", pad=0.2)
            fig.colorbar(im, cax=cax, label="Effective Polarization")
            plt.tight_layout(rect=[0, 0, 1, 0.96])
            self._save_fig_to_sheet(fig, "MRB Heatmap")
            return self

        # ---------- counts ----------
        def add_counts_sheet(self, raw_counts: Dict[str, Dict[str, int]], sheet_name="Counts"):
            ws = self.workbook.create_sheet(sheet_name)
            bits = sorted({b for d in raw_counts.values() for b in d})
            headers = ["Basis"] + bits
            for ci, h in enumerate(headers, 1):
                self._apply_header_style(ws.cell(1, ci, value=h))
            for ri, (lab, ct) in enumerate(sorted(raw_counts.items()), 2):
                ws.cell(ri, 1, lab).border = self._thin_border
                for ci, b in enumerate(bits, 2):
                    ws.cell(ri, ci, ct.get(b, 0)).border = self._thin_border
            self._auto_fit_columns(ws)
            return self
# ---------- probabilities ----------
        def add_probabilities_sheet(self, raw_counts: Dict[str, Dict[str, int]], sheet_name="Probabilities"):
            """Adds a worksheet showing measurement outcome probabilities."""
            ws = self.workbook.create_sheet(title=sheet_name)
            all_bits = sorted({b for d in raw_counts.values() for b in d})
            headers = ["Basis"] + [f"P({b})" for b in all_bits]
            for ci, h in enumerate(headers, 1):
                self._apply_header_style(ws.cell(1, ci, value=h))

            for ri, (basis, counts) in enumerate(sorted(raw_counts.items()), start=2):
                total = sum(counts.values())
                ws.cell(ri, 1, basis).border = self._thin_border
                for ci, b in enumerate(all_bits, start=2):
                    prob = counts.get(b, 0) / total if total > 0 else 0
                    cell = ws.cell(ri, ci, value=prob)
                    cell.border = self._thin_border
                    cell.number_format = "0.000000"
            self._auto_fit_columns(ws)
            return self
            

        # ---------- save ----------
        def save(self, filename):
            if not filename.endswith(".xlsx"):
                filename += ".xlsx"
            path = os.path.join(self.output_dir, filename)
            os.makedirs(os.path.dirname(path), exist_ok=True)
            self.workbook.save(path)
            print(f"[SUCCESS] Excel report saved to {path}")
            return path

        # ---------- util ----------
        def _auto_fit_columns(self, ws):
            for col in ws.columns:
                width = max((len(str(c.value)) for c in col if c.value), default=0) + 2
                ws.column_dimensions[get_column_letter(col[0].column)].width = width