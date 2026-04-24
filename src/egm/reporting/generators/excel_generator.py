# File Path: src/egm/reporting/generators/excel_generator.py
"""
A report generator that creates a detailed Excel (.xlsx) file.

This module contains the logic for creating multi-sheet Excel reports,
including summary tables, data tables, and embedded plots.
"""

import os
import io
import inspect
from typing import List, Dict, Any, Union

# Import visualization functions
from egm.reporting.visualizers import rb_plotter, heatmap_plotter, tomo_plotter

# Import schemas
from egm.schemas.results.base import BaseAnalysisResult
from egm.schemas.results.rb import RBAnalysisResult
from egm.schemas.results.tomo import QSTAnalysisResult, QPTAnalysisResult

# Type hint for any result
AnyAnalysisResult = Union[RBAnalysisResult, QSTAnalysisResult, QPTAnalysisResult]

try:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.drawing.image import Image as OpenpyxlImage
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
except ImportError:
    # If openpyxl is not installed, create a dummy class to avoid runtime errors.
    print("Warning: 'openpyxl' not found. Excel reporting features are disabled.")
    Workbook = None
    class ExcelGenerator:
        def __init__(self, *args, **kwargs):
            raise ImportError("'openpyxl' is required for ExcelGenerator.")

if Workbook:
    class ExcelGenerator:
        """
        Orchestrates the creation of a comprehensive Excel report from analysis results.
        """
        def __init__(self, output_dir: str = None, demo_name: str = None):
            if demo_name is None:
                # Auto-detect demo file name from the call stack
                demo_name = "unspecified_demo"
                for frame_info in inspect.stack():
                    if "demo_" in frame_info.filename:
                        demo_name = os.path.splitext(os.path.basename(frame_info.filename))[0]
                        break
            
            if output_dir is None:
                # Assume a project structure like 'src/egm/...' and place results in 'results_demo/'
                base_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "../../../.."))
                output_dir = os.path.join(base_path, "results_demo", demo_name)
            
            self.output_dir = os.path.abspath(output_dir)
            os.makedirs(self.output_dir, exist_ok=True)

            self.workbook = Workbook()
            self.workbook.remove(self.workbook.active)
            
            # Define cell styles
            self._header_font = Font(bold=True, color="FFFFFF")
            self._header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
            self._center_align = Alignment(horizontal="center", vertical="center")
            self._thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                                       top=Side(style='thin'), bottom=Side(style='thin'))

        def generate_report(
            self,
            analysis_results: List[AnyAnalysisResult],
            experiment_params: Dict[str, Any],
            raw_counts: Dict[str, Dict[str, int]] = None
        ) -> str:
            """
            Main method to generate and save the full report.

            Args:
                analysis_results: A list of result objects.
                experiment_params: A dictionary of experiment parameters.
                raw_counts: Optional dictionary of raw measurement counts.

            Returns:
                The file path to the saved Excel report.
            """
            self._create_summary_sheet(analysis_results, experiment_params)

            if raw_counts:
                self._add_counts_sheet(raw_counts, "Counts")
                self._add_probabilities_sheet(raw_counts, "Probabilities")

            for i, result in enumerate(analysis_results):
                self._add_plots_for_result(result, i + 1)
            
            filename = f"{experiment_params.get('name', 'report')}.xlsx"
            return self._save(filename)

        def _apply_header_style(self, cell):
            cell.font = self._header_font
            cell.fill = self._header_fill
            cell.alignment = self._center_align
            cell.border = self._thin_border

        def _auto_fit_columns(self, ws):
            for col in ws.columns:
                max_len = 0
                column = col[0].column_letter
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                adjusted_width = (max_len + 2)
                ws.column_dimensions[column].width = adjusted_width

        def _save_fig_to_sheet(self, fig, sheet_name: str):
            ws = self.workbook.create_sheet(title=sheet_name)
            with io.BytesIO() as buf:
                fig.savefig(buf, format="png", dpi=200, bbox_inches="tight")
                plt.close(fig)
                buf.seek(0)
                img = OpenpyxlImage(buf)
                ws.add_image(img, "A1")

        def _create_summary_sheet(self, results: List[AnyAnalysisResult], params: Dict[str, Any]):
            ws = self.workbook.create_sheet("Summary", 0)
            ws.cell(1, 1, "Experiment Parameters").font = Font(bold=True, size=14)
            row = 2
            for k, v in params.items():
                ws.cell(row, 1, k).font = Font(bold=True)
                ws.cell(row, 2, str(v))
                row += 1
            
            row += 1
            ws.cell(row, 1, "Key Results Summary").font = Font(bold=True, size=14)
            row += 1
            headers = ["Result ID", "Type", "Qubits", "Success", "Metric", "Value", "Error"]
            for i, h in enumerate(headers, 1):
                self._apply_header_style(ws.cell(row, i, h))

            for res in results:
                row += 1
                ws.cell(row, 1, str(res.result_id))
                ws.cell(row, 2, res.analysis_type)
                ws.cell(row, 3, str(res.qubits))
                ws.cell(row, 4, str(res.success))
                
                metric, value, error = "N/A", "N/A", "N/A"
                if isinstance(res, RBAnalysisResult):
                    metric, value, error = "EPC", res.epc, res.epc_error
                elif isinstance(res, QSTAnalysisResult):
                    metric, value = "Fidelity", res.fidelity
                elif isinstance(res, QPTAnalysisResult):
                    metric, value = "Proc. Fidelity", res.process_fidelity
                
                ws.cell(row, 5, metric)
                ws.cell(row, 6, value).number_format = '0.000E+00' if isinstance(value, float) and value < 0.01 else '0.00000'
                ws.cell(row, 7, error).number_format = '0.00E+00'

            self._auto_fit_columns(ws)

        def _add_plots_for_result(self, result: AnyAnalysisResult, index: int):
            """Dispatcher to add plots based on result type."""
            if isinstance(result, RBAnalysisResult):
                fig = rb_plotter.plot_rb_decay(result)
                self._save_fig_to_sheet(fig, f"RB Decay {index}")
                # Check for MRB data and add heatmap if present
                if hasattr(result, 'avg_polarizations'):
                    fig_hm = heatmap_plotter.plot_mrb_heatmap(result)
                    self._save_fig_to_sheet(fig_hm, f"MRB Heatmap {index}")
            elif isinstance(result, QSTAnalysisResult):
                fig = tomo_plotter.plot_density_matrix(result)
                self._save_fig_to_sheet(fig, f"Density Matrix {index}")
            elif isinstance(result, QPTAnalysisResult):
                fig = tomo_plotter.plot_chi_matrix(result)
                self._save_fig_to_sheet(fig, f"Chi Matrix {index}")

        def _add_counts_sheet(self, raw_counts: Dict[str, Dict[str, int]], sheet_name: str):
            ws = self.workbook.create_sheet(sheet_name)
            all_outcomes = sorted({b for d in raw_counts.values() for b in d})
            headers = ["Circuit/Basis"] + all_outcomes
            for ci, h in enumerate(headers, 1):
                self._apply_header_style(ws.cell(1, ci, value=h))
            
            for ri, (label, counts) in enumerate(sorted(raw_counts.items()), 2):
                ws.cell(ri, 1, label).border = self._thin_border
                for ci, outcome in enumerate(all_outcomes, 2):
                    ws.cell(ri, ci, counts.get(outcome, 0)).border = self._thin_border
            self._auto_fit_columns(ws)

        def _add_probabilities_sheet(self, raw_counts: Dict[str, Dict[str, int]], sheet_name: str):
            ws = self.workbook.create_sheet(sheet_name)
            all_outcomes = sorted({b for d in raw_counts.values() for b in d})
            headers = ["Circuit/Basis"] + [f"P({b})" for b in all_outcomes]
            for ci, h in enumerate(headers, 1):
                self._apply_header_style(ws.cell(1, ci, value=h))

            for ri, (basis, counts) in enumerate(sorted(raw_counts.items()), 2):
                total_shots = sum(counts.values())
                ws.cell(ri, 1, basis).border = self._thin_border
                for ci, outcome in enumerate(all_outcomes, 2):
                    prob = counts.get(outcome, 0) / total_shots if total_shots > 0 else 0
                    cell = ws.cell(ri, ci, value=prob)
                    cell.border = self._thin_border
                    cell.number_format = "0.000000"
            self._auto_fit_columns(ws)

        def _save(self, filename: str) -> str:
            path = os.path.join(self.output_dir, filename)
            self.workbook.save(path)
            print(f"[SUCCESS] Excel report saved to: {path}")
            return path